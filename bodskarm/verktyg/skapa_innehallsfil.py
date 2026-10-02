"""
Skapar data/Bodskarm_Innehall.xlsx – Excel-filen där du fyller i allt som
bodskärmen visar utöver 4-veckorsplaneringen.

Kör bara om du vill börja om från en tom mall (befintlig fil skrivs över!):

    python3 verktyg/skapa_innehallsfil.py            # mall med exempelrader
    python3 verktyg/skapa_innehallsfil.py --tom      # mall utan exempelrader
"""
import datetime
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

UT = os.path.join(os.path.dirname(__file__), "..", "data", "Bodskarm_Innehall.xlsx")
TOM = "--tom" in sys.argv

MORK = "13233B"
ACCENT = "3DA9FC"
FONT = "Calibri"

idag = datetime.date.today()
man = idag - datetime.timedelta(days=idag.weekday())


def dag(n):
    return idag + datetime.timedelta(days=n)


def vardag(n):
    """n:te vardagen från idag (0 = idag, eller nästa vardag om helg)."""
    d = idag
    while d.weekday() >= 5:
        d += datetime.timedelta(days=1)
    steg = 0
    while steg < n:
        d += datetime.timedelta(days=1)
        if d.weekday() < 5:
            steg += 1
    return d


def blad(wb, namn, titel, hjalp, kolumner, rader, bredder, tabellnamn, datumkol=(), tidkol=(), listor=None, ny=False):
    ws = wb.active if not ny else wb.create_sheet(namn)
    ws.title = namn
    ws.sheet_view.showGridLines = False
    ws["A1"] = titel
    ws["A1"].font = Font(name=FONT, size=18, bold=True, color=MORK)
    ws["A2"] = hjalp
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="5B6B85")
    ws["A2"].alignment = Alignment(wrap_text=False)
    ws.row_dimensions[1].height = 30

    for i, k in enumerate(kolumner, start=1):
        ws.cell(row=3, column=i, value=k)
    data = rader if rader else [[None] * len(kolumner)]
    for r, rad in enumerate(data, start=4):
        for c, v in enumerate(rad, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name=FONT, size=11)
            kol = kolumner[c - 1]
            if kol in datumkol:
                cell.number_format = "yyyy-mm-dd"
            if kol in tidkol:
                cell.number_format = "hh:mm"
    sista = 3 + len(data)
    ref = "A3:%s%d" % (chr(64 + len(kolumner)), sista)
    t = Table(displayName=tabellnamn, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for i, b in enumerate(bredder, start=1):
        ws.column_dimensions[chr(64 + i)].width = b
    ws.freeze_panes = "A4"

    # Formatera ett gäng tomma rader under tabellen så att nya rader ser rätt ut
    # när tabellen växer.
    for c, kol in enumerate(kolumner, start=1):
        bokstav = chr(64 + c)
        if kol in datumkol:
            for r in range(sista + 1, sista + 200):
                ws["%s%d" % (bokstav, r)].number_format = "yyyy-mm-dd"
        if kol in tidkol:
            for r in range(sista + 1, sista + 200):
                ws["%s%d" % (bokstav, r)].number_format = "hh:mm"

    for kol, varden in (listor or {}).items():
        c = kolumner.index(kol) + 1
        bokstav = chr(64 + c)
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(varden), allow_blank=True)
        dv.add("%s4:%s500" % (bokstav, bokstav))
        ws.add_data_validation(dv)
    return ws


def main():
    wb = Workbook()

    # ---------------------------------------------------------- Läs mig --
    ws = wb.active
    ws.title = "Läs mig"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    rader = [
        ("BODSKÄRM – innehåll", Font(name=FONT, size=22, bold=True, color=MORK)),
        ("Allt du skriver här visas på TV:n i boden. Spara filen – skärmen läser in den igen inom ett par minuter.", None),
        ("", None),
        ("Så här fungerar det", Font(name=FONT, size=14, bold=True, color=ACCENT)),
        ("• Tidplanen hämtas automatiskt från din 4-veckorsplanering. Den rör du som vanligt.", None),
        ("• I den här filen lägger du in det som inte finns i planeringen: meddelanden, hinder, leveranser, bilder och kontakter.", None),
        ("• Varje flik är en tabell. Lägg till en rad genom att skriva direkt under sista raden – tabellen växer av sig själv.", None),
        ("• Rader med passerat 'Visa till'-datum försvinner automatiskt från skärmen. Du behöver inte radera dem.", None),
        ("• Ändra inte rubrikerna (rad 3) – skärmen hittar kolumnerna via dem. Bladens namn ska också vara kvar.", None),
        ("", None),
        ("Flikarna", Font(name=FONT, size=14, bold=True, color=ACCENT)),
        ("Inställningar – projektnamn, väderort, vindgräns för lyft, säkerhetsinformation m.m.", None),
        ("Vyer – vilka sidor som visas, i vilken ordning och hur länge (sekunder).", None),
        ("Meddelanden – information till alla. Typ 'Säkerhet' och 'Viktigt' sorteras först. 'Rullande text = Ja' visar den även i textremsan längst ner.", None),
        ("Hinder & risker – öppna hinder visas. Sätt Status till 'Löst' när det är klart.", None),
        ("Leveranser – leveranser och kranlyft de närmaste 14 dagarna. Kranlyft varnas om vädret blåser över vindgränsen.", None),
        ("Bilder – skärmbilder från Trimble Connect, APD-plan, ritningar. Lägg bildfilen i mappen data\\bilder och skriv filnamnet här.", None),
        ("Kontakter – telefonlista som visas tillsammans med nödinformationen.", None),
        ("", None),
        ("Bilder från Trimble Connect", Font(name=FONT, size=14, bold=True, color=ACCENT)),
        ("1. Öppna modellen i Trimble Connect och ställ in vyn (gärna en sparad vy så att du hittar tillbaka).", None),
        ("2. Ta en skärmbild: Win + Shift + S, eller viewerns egen skärmbildsfunktion. Spara som PNG/JPG i data\\bilder.", None),
        ("3. Kopiera länken till vyn/modellen och klistra in i kolumnen 'Länk' – då visas en QR-kod som öppnar den i mobilen.", None),
        ("4. Tips: använd samma filnamn när du uppdaterar en vy, så behöver du inte ändra något här.", None),
    ]
    for i, (text, font) in enumerate(rader, start=1):
        c = ws.cell(row=i + 1, column=2, value=text)
        c.font = font or Font(name=FONT, size=12)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # ---------------------------------------------------- Inställningar --
    inst = [
        ["Projektnamn", "Nytt sovringsverk", "Stor rubrik uppe till vänster."],
        ["Underrubrik", "NSV Bygg · Hus del 2", "Liten text under projektnamnet."],
        ["Ort", "", "Visas i vädervyn, t.ex. 'Kiruna'."],
        ["Latitud", None, "Byggplatsens position för väderprognosen. Högerklicka på platsen i Google Maps så visas t.ex. 67.8558, 20.2253. Tom = ingen vädervy."],
        ["Longitud", None, "Andra talet från Google Maps."],
        ["Vindgräns lyft (m/s)", 12, "Varning visas när vindbyarna under arbetstid (06–18) når hit. Kontrollera mot kranens/lyftplanens gräns."],
        ["Standard visningstid (sek)", 20, "Används för vyer där 'Sekunder' är tomt på fliken Vyer."],
        ["Senaste olycka med frånvaro (datum)", None, "Visar 'X dagar utan olycka'. Tomt = visas inte."],
        ["Säkerhetsbudskap", "Rätt skyddsutrustning – varje dag, hela dagen." if not TOM else "", "Kort budskap i säkerhetsrutan på meddelandesidan."],
        ["Återsamlingsplats", "" if TOM else "EXEMPEL: Parkeringen vid grind 1", "Visas på kontaktsidan."],
        ["Hjärtstartare", "" if TOM else "EXEMPEL: Bod 3, vid entrén", ""],
        ["Första hjälpen", "" if TOM else "EXEMPEL: Bod 1 och bod 3", ""],
        ["Ögondusch", "", ""],
        ["Brandsläckare", "", ""],
        ["Adress för utryckning", "" if TOM else "EXEMPEL: Gatuadress 1, ort – infart via grind 1", "Det ambulans/räddningstjänst ska få höra."],
    ]
    blad(wb, "Inställningar", "Inställningar", "Fyll i kolumnen Värde. Ändra inte texterna i kolumnen Inställning.",
         ["Inställning", "Värde", "Förklaring"], inst, [36, 46, 110], "tblInstallningar", ny=True)

    # ------------------------------------------------------------- Vyer --
    vyer = [
        ["laget", "Läget just nu", "Ja", 20, 1, "Nyckeltal och framdrift per område."],
        ["vecka", "Denna vecka", "Ja", 25, 2, "Tidplan måndag–söndag för innevarande vecka."],
        ["fyraveckor", "4 veckor framåt", "Ja", 25, 3, "Tidplan för fyra veckor. Delas upp i flera sidor vid behov."],
        ["kommande", "Startar & ska bli klart", "Ja", 20, 4, "Vad som startar och ska bli klart de närmaste 14 dagarna."],
        ["bilder", "Modell & ritningar", "Ja", 20, 5, "En sida per bild på fliken Bilder."],
        ["leveranser", "Leveranser & lyft", "Ja", 20, 6, "Från fliken Leveranser."],
        ["hinder", "Hinder & risker", "Ja", 20, 7, "Öppna hinder från fliken Hinder & risker."],
        ["vader", "Väder", "Ja", 15, 8, "Kräver Latitud/Longitud under Inställningar."],
        ["meddelanden", "Meddelanden & säkerhet", "Ja", 20, 9, "Från fliken Meddelanden."],
        ["kontakter", "Kontakter & nödläge", "Ja", 15, 10, "Från fliken Kontakter + nödinfo under Inställningar."],
    ]
    blad(wb, "Vyer", "Vyer", "Visa = Nej stänger av en vy. Sekunder = visningstid per sida. Ordning = turordning. Ändra inte Vy-id.",
         ["Vy-id", "Namn", "Visa", "Sekunder", "Ordning", "Beskrivning"], vyer, [16, 30, 9, 11, 10, 80], "tblVyer",
         listor={"Visa": ["Ja", "Nej"]}, ny=True)

    # ------------------------------------------------------ Meddelanden --
    medd = [] if TOM else [
        ["Lyft över gångväg", "EXEMPEL: Mellan 07–10 på tisdag lyfts stomelement vid 744 Fläkthus. Gångvägen längs fasaden stängs.", "Säkerhet", dag(-2), dag(10), "Ja"],
        ["Skyddsrond fredag 13:00", "EXEMPEL: Samling vid bod 1. Arbetsledare och skyddsombud deltar.", "Viktigt", dag(-1), dag(14), "Ja"],
        ["Ny parkering för leveranser", "EXEMPEL: Leveransbilar hänvisas till ytan norr om 742. Se APD-planen.", "Info", dag(-3), dag(21), "Nej"],
        ["Fika fredag", "EXEMPEL: Gemensam fika i bod 2 kl 09:00.", "Info", dag(-1), dag(7), "Nej"],
    ]
    blad(wb, "Meddelanden", "Meddelanden", "Typ: Säkerhet / Viktigt / Info. Visa från/till styr när meddelandet syns (tomt = alltid).",
         ["Rubrik", "Text", "Typ", "Visa från", "Visa till", "Rullande text"], medd, [36, 80, 12, 13, 13, 14],
         "tblMeddelanden", datumkol=("Visa från", "Visa till"),
         listor={"Typ": ["Säkerhet", "Viktigt", "Info"], "Rullande text": ["Ja", "Nej"]}, ny=True)

    # --------------------------------------------------- Hinder & risker --
    hinder = [] if TOM else [
        ["742 Sikthall", "EXEMPEL: Armeringsritningar linje K rev C saknas", "Gjutning K10–K14 kan inte starta", "Påminna konstruktör, besked senast onsdag", "Arbetsledare bygg", vardag(2), "Öppen", "Hög"],
        ["744 Fläkthus", "EXEMPEL: Mobilkran ej bokad v." + str(dag(7).isocalendar()[1]), "Stomresning försenas", "Boka kran via inköp", "Platschef", vardag(3), "Pågår", "Medel"],
        ["745 Förtjockarhus", "EXEMPEL: Vatten i schakt efter regn", "Risk för urspolning av fyllning", "Pumpar på plats, kontroll varje morgon", "Arbetsledare mark", vardag(5), "Pågår", "Medel"],
        ["Allmänt", "EXEMPEL: Belysning gångväg vid bod 4 ur funktion", "Halk- och snubbelrisk i mörker", "Elektriker beställd", "Skyddsombud", vardag(1), "Öppen", "Låg"],
    ]
    blad(wb, "Hinder & risker", "Hinder & risker", "Status: Öppen / Pågår / Löst (lösta visas inte). Allvarlighet: Hög / Medel / Låg.",
         ["Område", "Beskrivning", "Konsekvens", "Åtgärd", "Ansvarig", "Klart senast", "Status", "Allvarlighet"], hinder,
         [20, 46, 38, 42, 20, 13, 11, 13], "tblHinder", datumkol=("Klart senast",),
         listor={"Status": ["Öppen", "Pågår", "Löst"], "Allvarlighet": ["Hög", "Medel", "Låg"]}, ny=True)

    # -------------------------------------------------------- Leveranser --
    lev = [] if TOM else [
        [vardag(0), datetime.time(7, 0), "EXEMPEL: Armering linje K", "Leverantör AB", "742 Sikthall, upplag norr", "Nej", "Arbetsledare bygg", "Lossas med truck"],
        [vardag(0), datetime.time(13, 30), "EXEMPEL: Betong 24 m³ C30/37", "Betongleverantör", "744 Fläkthus", "Nej", "Arbetsledare bygg", "Pump på plats 13:00"],
        [vardag(1), datetime.time(6, 30), "EXEMPEL: Stomelement etapp 2 (6 st)", "Prefableverantör", "744 Fläkthus", "Ja", "Platschef", "Mobilkran 60 t"],
        [vardag(2), datetime.time(9, 0), "EXEMPEL: Gjutformar", "Formleverantör", "745 Förtjockarhus", "Nej", "Arbetsledare mark", ""],
        [vardag(4), datetime.time(7, 0), "EXEMPEL: Stålpelare", "Stålleverantör", "742 Sikthall", "Ja", "Arbetsledare bygg", "Tornkran"],
    ]
    blad(wb, "Leveranser", "Leveranser", "Datum och tid för leveranser och lyft. Kranlyft = Ja ger vindvarning om det blåser.",
         ["Datum", "Tid", "Vad", "Leverantör", "Plats", "Kranlyft", "Mottagare", "Kommentar"], lev,
         [13, 9, 40, 24, 30, 10, 22, 34], "tblLeveranser", datumkol=("Datum",), tidkol=("Tid",),
         listor={"Kranlyft": ["Ja", "Nej"]}, ny=True)

    # ------------------------------------------------------------ Bilder --
    bilder = [] if TOM else [
        [1, "EXEMPEL: 744 Fläkthus – stomme etapp 2", "Ersätt med en skärmbild från Trimble Connect. Skriv gärna vad som ska göras i vyn.",
         "exempel_trimble_vy.svg", "https://web.connect.trimble.com/", "Trimble Connect", "Ja", None, None],
        [2, "EXEMPEL: APD-plan", "Uppställningsytor, kranplacering, gångvägar och återsamlingsplats.",
         "exempel_apd_plan.svg", "", "APD-plan", "Ja", None, None],
    ]
    blad(wb, "Bilder", "Bilder", "Lägg bildfilen (PNG/JPG) i mappen data\\bilder och skriv filnamnet i kolumnen Bildfil. Länk = QR-kod på skärmen.",
         ["Ordning", "Rubrik", "Beskrivning", "Bildfil", "Länk", "Kategori", "Visa", "Visa från", "Visa till"], bilder,
         [9, 40, 56, 28, 40, 18, 8, 13, 13], "tblBilder", datumkol=("Visa från", "Visa till"),
         listor={"Visa": ["Ja", "Nej"], "Kategori": ["Trimble Connect", "APD-plan", "Ritning", "Foto", "Övrigt"]}, ny=True)

    # --------------------------------------------------------- Kontakter --
    kontakter = [] if TOM else [
        ["Platschef", "EXEMPEL Namn", "070-000 00 01", "", "Ja"],
        ["Arbetsledare bygg", "EXEMPEL Namn", "070-000 00 02", "", "Ja"],
        ["Arbetsledare mark", "EXEMPEL Namn", "070-000 00 03", "", "Ja"],
        ["Skyddsombud", "EXEMPEL Namn", "070-000 00 04", "", "Ja"],
        ["BAS-U", "EXEMPEL Namn", "070-000 00 05", "", "Ja"],
        ["Beställarens kontakt", "EXEMPEL Namn", "070-000 00 06", "", "Ja"],
    ]
    blad(wb, "Kontakter", "Kontakter", "Visas på sidan Kontakter & nödläge (max 12 st). Visa = Nej döljer en rad.",
         ["Roll", "Namn", "Telefon", "Företag", "Visa"], kontakter, [28, 30, 20, 26, 8], "tblKontakter",
         listor={"Visa": ["Ja", "Nej"]}, ny=True)

    for ws in wb.worksheets[1:]:
        ws.sheet_properties.tabColor = ACCENT
    wb.active = 0
    os.makedirs(os.path.dirname(UT), exist_ok=True)
    wb.save(UT)
    print("Skapade", os.path.abspath(UT))


if __name__ == "__main__":
    main()

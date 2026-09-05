"""
Extraherar aktivitetsrader ur de fem områdesbladen (742, 744, 745, 775, 741)
i det ursprungliga NSV-produktionsunderlaget och skriver ut en ren,
strukturerad JSON-fil som sedan används för att bygga NSV Production Cockpit.

Körs en gång vid bygget. Kan köras om senare om källbladen ändras
väsentligt (nya aktiviteter tillagda/borttagna) - se README i AI ASSISTANT-
bladet för hur man uppdaterar 6-WEEK PLAN manuellt vid mindre ändringar.
"""
import json
import datetime
import openpyxl

SRC = "source/NSV_Bygg_4veckorsplanering_ORIGINAL.xlsm"

AREA_SHEETS = [
    ("741 - SEKTIONSFICKOR", "SEK"),
    ("742 - SIKTHALL", "SIK"),
    ("744 - Fläkthus", "FLA"),
    ("745 - FÖRTJOCKARHUS", "FOR"),
    ("775 - VATTENRESERVOAR", "VAT"),
]

# Kolumnlayout (identisk i alla områdesblad):
# A=V./Vecka  B=- (Del/Linje)  C=Aktivitet  D=Typ  E=Datum start  F=Antal dagar
# G=Datum slut  H=DP  I=BLOCK(Disciplin)  J=ÄTA  K=Framdrift %  L=Helgarbete

def main():
    wb = openpyxl.load_workbook(SRC, data_only=True, keep_vba=True)
    activities = []
    for sheet_name, prefix in AREA_SHEETS:
        ws = wb[sheet_name]
        seq = 0
        for r in range(5, ws.max_row + 1):
            start = ws.cell(row=r, column=5).value
            aktivitet = ws.cell(row=r, column=3).value
            if not isinstance(start, datetime.datetime) or not aktivitet:
                continue
            seq += 1
            act_id = f"{prefix}-{seq:04d}"
            slut = ws.cell(row=r, column=7).value
            antal_dagar = ws.cell(row=r, column=6).value
            record = {
                "id": act_id,
                "omrade": sheet_name,
                "del_linje": ws.cell(row=r, column=2).value or "",
                "aktivitet": str(aktivitet).strip(),
                "typ": ws.cell(row=r, column=4).value or "",
                "datum_start": start.strftime("%Y-%m-%d"),
                "antal_dagar": antal_dagar if isinstance(antal_dagar, (int, float)) else 0,
                "datum_slut": slut.strftime("%Y-%m-%d") if isinstance(slut, datetime.datetime) else start.strftime("%Y-%m-%d"),
                "dp": ws.cell(row=r, column=8).value or "",
                "disciplin": ws.cell(row=r, column=9).value or "",
                "ata": ws.cell(row=r, column=10).value or "Nej",
                "framdrift_kalla": ws.cell(row=r, column=11).value if isinstance(ws.cell(row=r, column=11).value, (int, float)) else 0,
                "helgarbete": ws.cell(row=r, column=12).value or "Nej",
                "kalla_rad": r,
            }
            activities.append(record)
        print(f"{sheet_name}: {seq} aktiviteter")

    print(f"TOTALT: {len(activities)} aktiviteter")

    with open("scripts/activities.json", "w", encoding="utf-8") as f:
        json.dump(activities, f, ensure_ascii=False, indent=1)

    # Sammanställ unika värden för dropdown-listor
    uniq = {
        "disciplin": sorted({a["disciplin"] for a in activities if a["disciplin"]}),
        "typ": sorted({a["typ"] for a in activities if a["typ"]}),
        "dp": sorted({a["dp"] for a in activities if a["dp"]}),
        "omrade": sorted({a["omrade"] for a in activities}),
    }
    print(json.dumps(uniq, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
